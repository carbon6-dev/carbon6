"""c6-render: carbon6's render kit (docs/docs.md §10).

pdf     <html> <out.pdf>            an HTML file to PDF (its @page sizes it)
slides  <html> <dir> <first> <last> a PNG per slide (slide-NN.png), 1280 wide
docx    <md> <out.docx>             Markdown to Word (pandoc)
odt     <md> <out.odt>              Markdown to OpenDocument text (pandoc)
pptx    <deck.json> <out.pptx>      a deck to PowerPoint, text editable
odp     <deck.json> <out.odp>       a deck to OpenDocument presentation
import  <file> <out.md>             docx, odt, pptx or md to carbon6 Markdown
split   <deck.html> <out.json>      an HTML deck (the decks app's): its styles
                                    and each <section class="slide"> as HTML

deck.json: {"title", "size": "16:9"|"4:3", "html": "deck.html", "slides":
[{"layout", "hidden", "notes", "picture": bool, "regions": {name: markdown}}]}
(pictures are the slides of deck.html, in order, hidden ones left out).
"""

import glob
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile


def fail(msg):
    sys.exit(msg)


def pdf(html, out):
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    profile = tempfile.mkdtemp(dir="/tmp")
    cmd = [
        "chromium", "--headless", "--no-sandbox", "--disable-gpu", "--disable-dev-shm-usage",
        "--no-first-run", "--user-data-dir=" + profile, "--crash-dumps-dir=/tmp",
        "--no-pdf-header-footer", "--virtual-time-budget=5000", "--run-all-compositor-stages-before-draw",
        "--print-to-pdf=" + os.path.abspath(out), "file://" + os.path.abspath(html),
    ]
    r = subprocess.run(cmd, capture_output=True, text=True)
    shutil.rmtree(profile, ignore_errors=True)
    if r.returncode != 0 or not os.path.exists(out):
        fail("chromium could not print %s:\n%s" % (html, r.stderr[-4000:]))


def pages(pdf_path, out_dir, first=0, last=0, width=1280):
    """One PNG per page, slide-NN.png; returns their paths in order."""
    os.makedirs(out_dir, exist_ok=True)
    tmp = tempfile.mkdtemp(dir="/tmp")
    cmd = ["pdftoppm", "-q", "-png", "-scale-to", str(width)]
    if first:
        cmd += ["-f", str(first), "-l", str(last or first)]
    subprocess.run(cmd + [pdf_path, os.path.join(tmp, "p")], check=True)
    out = []
    for p in sorted(glob.glob(os.path.join(tmp, "p-*.png")), key=lambda p: int(p.rsplit("-", 1)[1][:-4])):
        n = int(p.rsplit("-", 1)[1][:-4])
        dest = os.path.join(out_dir, "slide-%02d.png" % n)
        shutil.move(p, dest)
        out.append(dest)
    shutil.rmtree(tmp, ignore_errors=True)
    return out


def pandoc(src, out, to):
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    r = subprocess.run(["pandoc", "-f", "gfm", "-t", to, src, "-o", out, "--resource-path", os.path.dirname(os.path.abspath(src))],
                       capture_output=True, text=True)
    if r.returncode != 0:
        fail("pandoc: " + r.stderr[-4000:])


# --- Decks to PowerPoint -------------------------------------------------------

def plain(md):
    """Markdown inline marks taken off: **b**, *i*, `c`, [t](u)."""
    md = re.sub(r"\[([^\]]*)\]\([^)]*\)", r"\1", md)
    md = re.sub(r"(\*\*|__)(.+?)\1", r"\2", md)
    md = re.sub(r"(\*|_)(.+?)\1", r"\2", md)
    return md.replace("`", "")


def lines(md):
    """A region's Markdown as (level, text, bullet) lines."""
    out = []
    for raw in md.splitlines():
        if not raw.strip() or raw.strip().startswith("```") or set(raw.strip()) <= set("-|: "):
            continue
        indent = len(raw) - len(raw.lstrip())
        t = raw.strip()
        bullet = False
        m = re.match(r"^([-*+]|\d+[.)])\s+(\[[ xX]\]\s+)?(.*)$", t)
        if m:
            t, bullet = m.group(3), True
        t = re.sub(r"^#+\s*", "", t)
        t = re.sub(r"^>\s*", "", t)
        if t.startswith("|"):
            t = " · ".join(c.strip() for c in t.strip("|").split("|"))
        out.append((min(indent // 2, 4), plain(t), bullet))
    return out


def fill(frame, md, size=None):
    frame.clear()
    first = True
    for level, text, _bullet in lines(md):
        p = frame.paragraphs[0] if first else frame.add_paragraph()
        first = False
        p.text = text
        p.level = level
        if size:
            for r in p.runs:
                r.font.size = size


def pptx(deck_json, out):
    from pptx import Presentation
    from pptx.util import Emu, Pt
    d = json.load(open(deck_json))
    folder = os.path.dirname(os.path.abspath(deck_json))
    shown = [s for s in d["slides"] if not s.get("hidden")]
    pictures = {}
    if any(s.get("picture") for s in shown):
        tmp_pdf = os.path.join(tempfile.mkdtemp(dir="/tmp"), "deck.pdf")
        pdf(os.path.join(folder, d.get("html", "deck.html")), tmp_pdf)
        imgs = pages(tmp_pdf, tempfile.mkdtemp(dir="/tmp"), width=2560)
        pictures = {i: p for i, p in enumerate(imgs)}
    prs = Presentation()
    if d.get("size") == "4:3":
        prs.slide_width, prs.slide_height = Emu(9144000), Emu(6858000)
    else:
        prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    for i, s in enumerate(shown):
        r = s.get("regions", {})
        lay = s.get("layout", "body")
        if s.get("picture"):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            if i in pictures:
                slide.shapes.add_picture(pictures[i], 0, 0, width=prs.slide_width, height=prs.slide_height)
        elif lay in ("title", "section"):
            slide = prs.slides.add_slide(prs.slide_layouts[0 if lay == "title" else 2])
            fill(slide.shapes.title.text_frame, r.get("title", ""))
            if len(slide.placeholders) > 1:
                fill(slide.placeholders[1].text_frame, r.get("subtitle", ""))
        elif lay == "two-columns":
            slide = prs.slides.add_slide(prs.slide_layouts[3])
            fill(slide.shapes.title.text_frame, r.get("title", ""))
            fill(slide.placeholders[1].text_frame, r.get("left", ""))
            fill(slide.placeholders[2].text_frame, r.get("right", ""))
        else:
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            title = r.get("title", "")
            body = r.get("body", "")
            if lay == "big-number":
                body = "\n".join(x for x in (r.get("number", ""), r.get("caption", "")) if x)
            elif lay == "quote":
                title = title or r.get("quote", "")
                body = r.get("attribution", "") if r.get("title") is None else r.get("quote", "") + "\n" + r.get("attribution", "")
            elif lay in ("image", "chart", "diagram"):
                body = r.get("caption", "")
            fill(slide.shapes.title.text_frame, title)
            fill(slide.placeholders[1].text_frame, body, Pt(60) if lay == "big-number" else None)
        notes = s.get("notes") or ""
        if s.get("data"):
            notes = (notes + "\n\n" + s["data"]).strip()
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    prs.save(out)
    print("%d slides in %s" % (len(shown), out))


def odp(deck_json, out):
    tmp = tempfile.mkdtemp(dir="/tmp")
    p = os.path.join(tmp, "deck.pptx")
    pptx(deck_json, p)
    r = subprocess.run(["soffice", "--headless", "-env:UserInstallation=file:///tmp/lo", "--convert-to", "odp", "--outdir", tmp, p],
                       capture_output=True, text=True)
    made = os.path.join(tmp, "deck.odp")
    if not os.path.exists(made):
        fail("libreoffice could not convert: " + (r.stderr or r.stdout)[-4000:])
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    shutil.move(made, out)


# --- Files into documents ------------------------------------------------------

def import_pptx(path):
    from pptx import Presentation
    prs = Presentation(path)
    slides = []
    for n, s in enumerate(prs.slides):
        title = s.shapes.title.text.strip() if s.shapes.title is not None and s.shapes.title.has_text_frame else ""
        body, pictures = [], 0
        for sh in s.shapes:
            if s.shapes.title is not None and sh.shape_id == s.shapes.title.shape_id:
                continue
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    t = "".join(r.text for r in p.runs).strip()
                    if t:
                        body.append("  " * p.level + "- " + t)
            elif sh.shape_type == 13:
                pictures += 1
        notes = ""
        if s.has_notes_slide:
            notes = s.notes_slide.notes_text_frame.text.strip()
        layout = "title" if n == 0 and len(body) <= 1 else "body"
        md = "<!-- layout: %s -->\n" % layout
        md += ("# " if layout == "title" else "## ") + (title or "Slide %d" % (n + 1)) + "\n"
        if layout == "title" and body:
            md += "\n" + body[0][2:] + "\n"
        elif body:
            md += "\n" + "\n".join(body) + "\n"
        if pictures:
            md += "\n*(%d picture%s on the original slide)*\n" % (pictures, "" if pictures == 1 else "s")
        if notes:
            md += "\n<!-- notes: %s -->\n" % notes.replace("-->", "-- >")
        slides.append(md)
    return "\n---\n\n".join(slides)


def import_file(path, out):
    ext = os.path.splitext(path)[1].lower()
    if ext in (".md", ".markdown", ".txt"):
        text = open(path, encoding="utf-8", errors="replace").read()
    elif ext == ".pptx":
        text = import_pptx(path)
    elif ext in (".docx", ".odt", ".html", ".htm", ".rtf"):
        r = subprocess.run(["pandoc", path, "-t", "gfm-raw_html", "--wrap=none"], capture_output=True, text=True)
        if r.returncode != 0:
            fail("pandoc: " + r.stderr[-4000:])
        text = r.stdout
    else:
        fail("can't read %s files: Word, OpenDocument text, PowerPoint or Markdown" % ext)
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    open(out, "w", encoding="utf-8").write(text)


def split(html, out):
    """The decks app's deck.html: its <style>s, and each slide's HTML."""
    import lxml.html
    tree = lxml.html.parse(html).getroot()
    styles = "\n".join(s.text or "" for s in tree.iter("style"))
    sections = []
    for sec in tree.iter("section"):
        cls = (sec.get("class") or "").split()
        if "slide" in cls or not any("slide" in (x.get("class") or "") for x in tree.iter("section")):
            sections.append(lxml.html.tostring(sec, encoding="unicode", method="html"))
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    json.dump({"style": styles, "sections": sections}, open(out, "w"))
    print("%d slides" % len(sections))


def main():
    if len(sys.argv) < 4:
        fail(__doc__)
    what, src, out = sys.argv[1], sys.argv[2], sys.argv[3]
    if what == "pdf":
        pdf(src, out)
    elif what == "slides":
        first = int(sys.argv[4] or 0) if len(sys.argv) > 4 and sys.argv[4] else 0
        last = int(sys.argv[5] or 0) if len(sys.argv) > 5 and sys.argv[5] else 0
        tmp_pdf = os.path.join(tempfile.mkdtemp(dir="/tmp"), "deck.pdf")
        pdf(src, tmp_pdf)
        made = pages(tmp_pdf, out, first, last)
        print("%d slides" % len(made))
    elif what in ("docx", "odt"):
        pandoc(src, out, what)
    elif what == "pptx":
        pptx(src, out)
    elif what == "odp":
        odp(src, out)
    elif what == "import":
        import_file(src, out)
    elif what == "split":
        split(src, out)
    else:
        fail("unknown: " + what)


if __name__ == "__main__":
    main()
