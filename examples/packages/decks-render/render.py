"""decks-render: an HTML deck (one page per slide, sized by its @page rule)
to PDF, to one PNG per slide, or to PowerPoint (a picture per slide)."""

import glob
import os
import shutil
import subprocess
import sys
import tempfile


def pdf(html, out):
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    profile = tempfile.mkdtemp(dir="/tmp")
    cmd = [
        "chromium", "--headless", "--no-sandbox", "--disable-gpu", "--disable-dev-shm-usage",
        "--no-first-run", "--user-data-dir=" + profile, "--crash-dumps-dir=/tmp",
        "--no-pdf-header-footer", "--virtual-time-budget=5000", "--run-all-compositor-stages-before-draw",
        "--print-to-pdf=" + os.path.abspath(out), "file://" + os.path.abspath(html),
    ]
    r = subprocess.run(cmd, capture_output=True, text=True)
    shutil.rmtree(profile, ignore_errors=True)
    if r.returncode != 0 or not os.path.exists(out):
        sys.exit("chromium could not print %s:\n%s" % (html, r.stderr[-4000:]))


def pages(pdf_path, out_dir, first=None, last=None, width=1280):
    """One PNG per page, slide-01.png …; returns their paths."""
    os.makedirs(out_dir, exist_ok=True)
    tmp = tempfile.mkdtemp(dir="/tmp")
    cmd = ["pdftoppm", "-q", "-png", "-scale-to", str(width)]
    if first:
        cmd += ["-f", str(first), "-l", str(last or first)]
    subprocess.run(cmd + [pdf_path, os.path.join(tmp, "p")], check=True)
    out = []
    for p in sorted(glob.glob(os.path.join(tmp, "p-*.png")), key=lambda p: int(p.rsplit("-", 1)[1][:-4])):
        n = int(p.rsplit("-", 1)[1][:-4])
        dest = os.path.join(out_dir, "slide-%02d.png" % n)
        shutil.move(p, dest)
        out.append(dest)
    shutil.rmtree(tmp, ignore_errors=True)
    return out


def main():
    if len(sys.argv) < 4:
        sys.exit("usage: render.py pdf|slides|slide|pptx <deck.html> <out> [n]")
    what, html, out = sys.argv[1], sys.argv[2], sys.argv[3]
    if what == "pdf":
        pdf(html, out)
        print("printed %s" % out)
        return
    tmp_pdf = os.path.join(tempfile.mkdtemp(dir="/tmp"), "deck.pdf")
    pdf(html, tmp_pdf)
    if what == "slides":
        made = pages(tmp_pdf, out)
        print("%d slides: %s" % (len(made), ", ".join(made)))
    elif what == "slide":
        n = int(sys.argv[4])
        made = pages(tmp_pdf, out, n, n)
        if not made:
            sys.exit("the deck has no slide %d" % n)
        print(made[0])
    elif what == "pptx":
        from pptx import Presentation
        from pptx.util import Emu
        imgs = pages(tmp_pdf, tempfile.mkdtemp(dir="/tmp"), width=2560)
        prs = Presentation()
        prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
        for img in imgs:
            s = prs.slides.add_slide(prs.slide_layouts[6])
            s.shapes.add_picture(img, 0, 0, width=prs.slide_width, height=prs.slide_height)
        os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
        prs.save(out)
        print("%d slides in %s" % (len(imgs), out))
    else:
        sys.exit("unknown: " + what)


if __name__ == "__main__":
    main()
